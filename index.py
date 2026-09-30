import streamlit as st
import requests
import re
from fpdf import FPDF
from docx import Document
from io import BytesIO


# =========================================================
# PAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="centered"
)


# =========================================================
# TITLE
# =========================================================

st.title("⚖️ LegalEase")
st.subheader("AI-Powered Legal Document Generator")

st.write(
    "Generate professional legal documents using Artificial Intelligence."
)

st.markdown("""
<style>
.main-title {
    text-align: center;
    font-size: 42px;
    font-weight: bold;
    margin-bottom: 5px;
}

.subtitle {
    text-align: center;
    font-size: 18px;
    margin-bottom: 25px;
}

.stButton > button {
    width: 100%;
    border-radius: 10px;
    height: 48px;
    font-size: 16px;
    font-weight: bold;
}

div[data-testid="stDownloadButton"] button {
    width: 100%;
    border-radius: 10px;
}
.legal-header {
    display: flex;
    align-items: center;
    justify-content: center;
    gap: 15px;
    margin-bottom: 10px;
}

.logo {
    font-size: 48px;
}

.brand-name {
    font-size: 38px;
    font-weight: bold;
}

.brand-tagline {
    font-size: 16px;
}

.description {
    text-align: center;
    font-size: 16px;
    margin-bottom: 30px;
}
</style>
""", unsafe_allow_html=True)

# =========================================================
# DOCUMENT TYPE
# =========================================================

document_type = st.selectbox(
    "Select Document Type",
    [
        "Employment Contract",
        "Non-Disclosure Agreement",
        "Lease Agreement",
        "Service Agreement"
    ]
)


# =========================================================
# INPUT DETAILS
# =========================================================

st.subheader("Enter Document Details")


name = st.text_input(
    "Name",
    placeholder="Enter employee / party name"
)


company = st.text_input(
    "Company / Organization",
    placeholder="Enter company name"
)


job_role = st.text_input(
    "Job Role",
    placeholder="Example: Software Developer"
)


salary = st.text_input(
    "Salary",
    placeholder="Example: 30000 per month"
)


start_date = st.text_input(
    "Start / Effective Date",
    placeholder="Example: 1 October 2026"
)


# =========================================================
# GENERATE DOCUMENT
# =========================================================

if st.button("Generate Document"):

    # -----------------------------------------------------
    # BASIC VALIDATION
    # -----------------------------------------------------

    if not name.strip() or not company.strip():

        st.warning(
            "Please enter at least the Name and Company."
        )

    else:

        # -------------------------------------------------
        # HANDLE EMPTY DETAILS
        # -------------------------------------------------

        safe_name = name.strip()
        safe_company = company.strip()

        safe_job_role = (
            job_role.strip()
            if job_role.strip()
            else "Not specified"
        )

        safe_salary = (
            salary.strip()
            if salary.strip()
            else "Not specified"
        )

        safe_start_date = (
            start_date.strip()
            if start_date.strip()
            else "Not specified"
        )


        # -------------------------------------------------
        # DETAILS SENT TO BACKEND
        # -------------------------------------------------

        details = f"""
Name: {safe_name}
Company: {safe_company}
Job Role: {safe_job_role}
Salary: {safe_salary}
Start Date: {safe_start_date}
"""


        # -------------------------------------------------
        # CALL BACKEND
        # -------------------------------------------------

        with st.spinner("Generating your document..."):

            try:

                response = requests.post(
                    "http://127.0.0.1:8000/generate",

                    json={
                        "document_type": document_type,
                        "details": details
                    },

                    timeout=120
                )


                # =================================================
                # SUCCESS
                # =================================================

                if response.status_code == 200:

                    result = response.json()

                    generated_document = result.get(
                        "generated_document",
                        ""
                    )


                    if generated_document:

                        # -----------------------------------------
                        # CLEAN AI FORMATTING
                        # -----------------------------------------

                        clean_text = generated_document

                        # Remove markdown bold
                        clean_text = clean_text.replace(
                            "**",
                            ""
                        )

                        # Remove markdown headings
                        clean_text = re.sub(
                            r"^#{1,6}\s*",
                            "",
                            clean_text,
                            flags=re.MULTILINE
                        )

                        # Convert <br>
                        clean_text = re.sub(
                            r"<br\s*/?>",
                            "\n",
                            clean_text,
                            flags=re.IGNORECASE
                        )

                        # Replace special characters
                        clean_text = clean_text.replace(
                            "—",
                            "-"
                        )

                        clean_text = clean_text.replace(
                            "–",
                            "-"
                        )

                        clean_text = clean_text.replace(
                            "“",
                            '"'
                        )

                        clean_text = clean_text.replace(
                            "”",
                            '"'
                        )

                        clean_text = clean_text.replace(
                            "‘",
                            "'"
                        )

                        clean_text = clean_text.replace(
                            "’",
                            "'"
                        )


                        # -----------------------------------------
                        # SAVE DOCUMENT
                        # -----------------------------------------

                        st.session_state["generated_document"] = clean_text


                        # -----------------------------------------
                        # SUCCESS MESSAGE
                        # -----------------------------------------

                        st.success(
                            "Document generated successfully!"
                        )


                    else:

                        st.error(
                            "Backend returned an empty document."
                        )


                # =================================================
                # BACKEND ERROR
                # =================================================

                else:

                    st.error(
                        f"Backend Error: {response.status_code}"
                    )

                    st.write(
                        "Backend Response:"
                    )

                    st.code(
                        response.text
                    )


            # =====================================================
            # CONNECTION ERROR
            # =====================================================

            except requests.exceptions.ConnectionError:

                st.error(
                    "Backend is not running."
                )

                st.info(
                    "Please start FastAPI first."
                )


            # =====================================================
            # TIMEOUT
            # =====================================================

            except requests.exceptions.Timeout:

                st.error(
                    "Request timed out. Please try again."
                )


            # =====================================================
            # OTHER ERROR
            # =====================================================

            except Exception as e:

                st.error(
                    f"Something went wrong: {e}"
                )


# =========================================================
# DOCUMENT PREVIEW
# =========================================================

if "generated_document" in st.session_state:

    generated_document = st.session_state[
        "generated_document"
    ]


    st.divider()

    st.subheader("📄 Document Preview")


    # ---------------------------------------------------------
    # EDIT DOCUMENT
    # ---------------------------------------------------------

    edited_document = st.text_area(
        "Edit your document",
        value=generated_document,
        height=600
    )


    # =========================================================
    # TXT DOWNLOAD
    # =========================================================

    txt_file = edited_document.encode(
        "utf-8"
    )


    st.download_button(
        label="📥 Download as TXT",
        data=txt_file,
        file_name=f"{document_type}.txt",
        mime="text/plain"
    )


    # =========================================================
    # PDF DOWNLOAD
    # =========================================================

    try:

        pdf = FPDF()

        pdf.set_auto_page_break(
            auto=True,
            margin=15
        )

        pdf.add_page()


        # Title
        pdf.set_font(
            "Arial",
            "B",
            16
        )

        pdf.cell(
            0,
            10,
            document_type,
            ln=True,
            align="C"
        )

        pdf.ln(8)


        # Body
        pdf.set_font(
            "Arial",
            "",
            11
        )


        for line in edited_document.split("\n"):

            line = line.strip()

            if line == "":

                pdf.ln(5)

            else:

                # Make text PDF-safe
                safe_line = line.encode(
                    "latin-1",
                    "replace"
                ).decode(
                    "latin-1"
                )

                pdf.multi_cell(
                    0,
                    7,
                    safe_line
                )


        # Create PDF bytes
        pdf_output = pdf.output(
            dest="S"
        )

        if isinstance(pdf_output, str):

            pdf_file = pdf_output.encode(
                "latin-1"
            )

        else:

            pdf_file = bytes(pdf_output)


        st.download_button(
            label="📄 Download as PDF",
            data=pdf_file,
            file_name=f"{document_type}.pdf",
            mime="application/pdf"
        )


    except Exception as e:

        st.error(
            f"PDF creation error: {e}"
        )


    # =========================================================
    # DOCX DOWNLOAD
    # =========================================================

    try:

        doc = Document()


        # -----------------------------------------------------
        # DOCX TITLE
        # -----------------------------------------------------

        title = doc.add_heading(
            document_type,
            level=1
        )

        title.alignment = 1


        # -----------------------------------------------------
        # DOCX CONTENT
        # -----------------------------------------------------

        for line in edited_document.split("\n"):

            line = line.strip()

            if line == "":

                doc.add_paragraph("")

            else:

                doc.add_paragraph(line)


        # -----------------------------------------------------
        # SAVE DOCX TO MEMORY
        # -----------------------------------------------------

        docx_file = BytesIO()

        doc.save(
            docx_file
        )

        docx_file.seek(0)


        # -----------------------------------------------------
        # DOWNLOAD BUTTON
        # -----------------------------------------------------

        st.download_button(
            label="📝 Download as DOCX",
            data=docx_file,
            file_name=f"{document_type}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


    except Exception as e:

        st.error(
            f"DOCX creation error: {e}"
        )